import pptx
from pptx.dml.color import RGBColor

prs = pptx.Presentation(r'C:\Users\HP\Downloads\QueryPilot_Project_Review_Presentation.pptx')

print("=== SLIDE 2 STYLES ===")
slide2 = prs.slides[1]
for s in slide2.shapes:
    if s.has_text_frame:
        for p in s.text_frame.paragraphs:
            if p.text.strip():
                font_name = p.font.name
                font_size = p.font.size.pt if p.font.size else None
                color = None
                if p.runs and p.runs[0].font.color and p.runs[0].font.color.type == 1:
                    color = p.runs[0].font.color.rgb
                print(f"Shape: {s.name} | text: {p.text[:40]} | font: {font_name} | size: {font_size} | color: {color}")

print("\n=== SLIDE 13 (Implementation Status) STYLES ===")
slide13 = prs.slides[12]
for s in slide13.shapes:
    print(f"Shape: {s.name}, type: {s.shape_type}, left: {s.left/914400:.2f}, width: {s.width/914400:.2f}")
    if s.has_text_frame:
        print(f"  text: '{s.text_frame.text.strip()}'")

print("\n=== SLIDE 6 (Table) STYLES ===")
slide6 = prs.slides[5]
for s in slide6.shapes:
    if s.has_table:
        t = s.table
        print(f"Table rows={len(t.rows)}, cols={len(t.columns)}")
        for r_idx, r in enumerate(t.rows):
            row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            print(f"  Row {r_idx}: {row_vals}")

print("\n=== SLIDE 12 (API) STYLES ===")
slide12 = prs.slides[11]
for s in slide12.shapes:
    if s.has_table:
        t = s.table
        print(f"Table rows={len(t.rows)}, cols={len(t.columns)}")
        for r_idx, r in enumerate(t.rows):
            row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            print(f"  Row {r_idx}: {row_vals}")

print("\n=== SLIDE 15 (Challenges) STYLES ===")
slide15 = prs.slides[14]
for s in slide15.shapes:
    if s.has_table:
        t = s.table
        print(f"Table rows={len(t.rows)}, cols={len(t.columns)}")
        for r_idx, r in enumerate(t.rows):
            row_vals = [c.text.strip().replace('\n', ' ') for c in r.cells]
            print(f"  Row {r_idx}: {row_vals}")
