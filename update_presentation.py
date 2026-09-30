import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# Color constants matching the QueryPilot presentation UI/UX design
COLOR_DARK_TEXT = RGBColor(0x1B, 0x2A, 0x2E)    # #1B2A2E Charcoal
COLOR_TEAL_HEADER = RGBColor(0x02, 0x80, 0x90)  # #028090 Deep Teal
COLOR_TEAL_ACCENT = RGBColor(0x00, 0xA8, 0x96)  # #00A896 Vibrant Teal
COLOR_TEAL_LIGHT = RGBColor(0x02, 0xC3, 0x9A)   # #02C39A Mint/Emerald
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)        # #FFFFFF
COLOR_MUTED = RGBColor(0x7A, 0x8F, 0x8C)        # #7A8F8C Gray-Teal

def format_cell(cell, text, is_header=False, font_size=12):
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for r in p.runs:
        r.font.name = "Calibri"
        r.font.size = Pt(12.5 if is_header else font_size)
        r.font.bold = is_header
        r.font.color.rgb = COLOR_WHITE if is_header else COLOR_DARK_TEXT
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_TEAL_HEADER if is_header else COLOR_WHITE

def set_paragraph_text(p, text, bold=False, color=COLOR_DARK_TEXT, size=13):
    p.text = text
    p.font.name = "Calibri"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color

def set_bullet_list(text_frame, items, size=12):
    text_frame.clear()
    for idx, item in enumerate(items):
        p = text_frame.add_paragraph() if idx > 0 else text_frame.paragraphs[0]
        p.text = item
        p.font.name = "Calibri"
        p.font.size = Pt(size)
        p.font.color.rgb = COLOR_DARK_TEXT
        p.space_after = Pt(6)

def update_presentation():
    src_path = r"C:\Users\HP\Downloads\QueryPilot_Project_Review_Presentation.pptx"
    prs = pptx.Presentation(src_path)
    print("Loaded presentation with", len(prs.slides), "slides")

    # ----------------------------------------------------
    # SLIDE 1 (Agenda): Update agenda items if needed
    # ----------------------------------------------------
    slide1 = prs.slides[0]
    for s in slide1.shapes:
        if s.name == "Text 8" and s.has_text_frame:
            agenda_items = [
                "Project Overview & Core Mission",
                "Requirements (Functional & Non-Functional)",
                "System Architecture (3-Tier & AI Performance Engine)",
                "Use Case & Database ER Diagrams (MongoDB 1:N)",
                "System Modules & Technology Stack",
                "UI Design, Frontend & Backend Engineering",
                "Database Integration & Real Express REST APIs",
                "Implementation Status & Telemetry HUD",
                "Live Application Screenshots & Real-World Challenges",
                "Strategic Roadmap & Official References"
            ]
            set_bullet_list(s.text_frame, agenda_items, size=11.5)

    # ----------------------------------------------------
    # SLIDE 2 (01 Requirements): Functional & Non-Functional
    # ----------------------------------------------------
    slide2 = prs.slides[1]
    for s in slide2.shapes:
        if s.name == "Text 7" and s.has_text_frame:
            functional_reqs = [
                "User authentication with JWT (Sign Up, Sign In, session persistence).",
                "Raw query input & execution plan ingestion (SQL / MongoDB query format).",
                "Automated performance diagnostics: Full-table scan detection, join cost analysis, and missing index flags.",
                "Query rewrite and indexing recommendations generated via the Analyzer engine.",
                "Historical analysis report archiving and client-side PDF export (exportPdf.js)."
            ]
            set_bullet_list(s.text_frame, functional_reqs, size=12)
        elif s.name == "Text 10" and s.has_text_frame:
            non_functional_reqs = [
                "Performance: Analysis turnaround under 1.5 seconds per query.",
                "Security: Secure password hashing using bcryptjs and route protection with Bearer tokens.",
                "Usability: Modern, responsive interface with code editors and metrics visualization.",
                "Reliability: MongoDB persistence with health-check monitoring and data isolation.",
                "Maintainability: Modular React components and decoupled Express service layers."
            ]
            set_bullet_list(s.text_frame, non_functional_reqs, size=12)

    # ----------------------------------------------------
    # SLIDE 3 (02 Architecture): Add AI Analyzer and PDF Engine
    # ----------------------------------------------------
    slide3 = prs.slides[2]
    for s in slide3.shapes:
        if s.name == "Text 7" and s.has_text_frame:
            client_items = [
                "React.js + Vite (Code Editor & HUD)",
                "Client PDF Generator (jsPDF + autotable)",
                "Session State & Local Storage Management"
            ]
            set_bullet_list(s.text_frame, client_items, size=11.5)
        elif s.name == "Text 10" and s.has_text_frame:
            s.text_frame.paragraphs[0].text = "APPLICATION & AI TIER"
            s.text_frame.paragraphs[0].font.name = "Calibri"
            s.text_frame.paragraphs[0].font.size = Pt(13)
            s.text_frame.paragraphs[0].font.bold = True
            s.text_frame.paragraphs[0].font.color.rgb = COLOR_TEAL_HEADER
        elif s.name == "Text 11" and s.has_text_frame:
            app_items = [
                "Node.js + Express.js REST Framework",
                "Security Middleware (JWT Auth & bcryptjs)",
                "AI Analyzer Service (analyzer.js): AST parsing, scan detection & index advice"
            ]
            set_bullet_list(s.text_frame, app_items, size=11.5)
        elif s.name == "Text 15" and s.has_text_frame:
            data_items = [
                "MongoDB NoSQL Cloud Cluster",
                "users & reports collections",
                "Unique & compound indexes for sub-second retrieval"
            ]
            set_bullet_list(s.text_frame, data_items, size=11.5)
        elif s.name == "Text 16" and s.has_text_frame:
            s.text_frame.text = "Data Flow: Client submits raw queries/execution plans via REST endpoints; AI Analyzer engine performs heuristic bottleneck detection and index advice; results persisted to MongoDB and exported as client-side PDF."
            p = s.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.size = Pt(11.5)
            p.font.color.rgb = COLOR_DARK_TEXT
        elif s.name == "Text 18" and s.has_text_frame:
            s.text_frame.text = "Key Architectural Enhancement: Decoupled AI diagnostic engine (analyzer.js) isolates compute-intensive AST query parsing and rule evaluation from HTTP handling."
            p = s.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_DARK_TEXT

    # ----------------------------------------------------
    # SLIDE 4 (03 Use Case Diagram): Insert Diagram & Description
    # ----------------------------------------------------
    slide4 = prs.slides[3]
    # Update right card texts
    for s in slide4.shapes:
        if s.name == "Text 8" and s.has_text_frame:
            s.text_frame.text = "Use Case Specification"
            p = s.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEAL_HEADER
        elif s.name == "Text 9" and s.has_text_frame:
            uc_items = [
                "Primary Actor: Database Developer / DBA / Registered User",
                "Core System Use Cases:",
                "• User Sign-Up & Secure Login (JWT token & session persistence)",
                "• Ingest Raw SQL Query & Execution Plan breakdown",
                "• Automated Performance Diagnostics: Detect full-table scans & joins",
                "• AI Recommendation Engine: Composite B-Tree indexes & rewrites",
                "• Query History Archiving: Persist & filter past runs in MongoDB",
                "• Client-Side PDF Generation: Instant download of DBA audit reports"
            ]
            set_bullet_list(s.text_frame, uc_items, size=11)

    # Insert use_case_diagram.png into left card area
    diag_img = r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\use_case_diagram.png"
    if os.path.exists(diag_img):
        slide4.shapes.add_picture(diag_img, Inches(1.25), Inches(1.95), Inches(6.50), Inches(4.40))
        print("Inserted use_case_diagram.png into Slide 4")

    # ----------------------------------------------------
    # SLIDE 5 (04 ER Diagram / Database Design): Insert Diagram & Details
    # ----------------------------------------------------
    slide5 = prs.slides[4]
    for s in slide5.shapes:
        if s.name == "Text 8" and s.has_text_frame:
            s.text_frame.text = "MongoDB Schema (1:N)"
            p = s.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEAL_HEADER
        elif s.name == "Text 9" and s.has_text_frame:
            er_items = [
                "1. User Entity (models/User.js):",
                "• _id: ObjectId (Primary Key)",
                "• name: String (User display name)",
                "• email: String (Unique index, normalized lowercase)",
                "• password: String (Hashed with bcryptjs)",
                "• createdAt: Timestamp (Account registration date)",
                "2. Report Entity (models/Report.js):",
                "• _id: ObjectId (Primary Key)",
                "• userId: ObjectId (Foreign Key -> User._id)",
                "• query: String (Raw SQL / Execution plan)",
                "• databaseType: String (MySQL, PostgreSQL, MongoDB)",
                "• executionTime / cost: Number/Object (Score metrics)",
                "• recommendations: Array/Text (AI rewrites & indexes)",
                "• createdAt: Timestamp (Compound indexed)"
            ]
            set_bullet_list(s.text_frame, er_items, size=10.5)

    er_img = r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\er_diagram.png"
    if os.path.exists(er_img):
        slide5.shapes.add_picture(er_img, Inches(1.25), Inches(1.95), Inches(6.50), Inches(4.40))
        print("Inserted er_diagram.png into Slide 5")

    # ----------------------------------------------------
    # SLIDE 6 (05 Module Description): Exact Replacement Table
    # ----------------------------------------------------
    slide6 = prs.slides[5]
    for s in slide6.shapes:
        if s.has_table:
            t = s.table
            modules_data = [
                ("Module", "Description", "Key Functions"),
                ("Authentication Module", "Secure user onboarding and session management.", "User registration, login, bcrypt password hashing, JWT token validation."),
                ("Query Diagnostic Engine", "Parses database queries and execution plans.", "Syntax parsing, missing index detection, scan-type identification, execution cost scoring."),
                ("AI Advisor & Recommendation", "Synthesizes optimized query structures and index strategies.", "Suggests composite indexes, query rewrites, schema indexing suggestions."),
                ("Report History & Export", "Persists diagnostic history and exports client deliverables.", "Save report to MongoDB, filter past analyses, generate downloadable PDF reports."),
                ("Analytics & Telemetry HUD", "Aggregates performance telemetry and historical benchmarks.", "Visual score gauges (0-100), risk tier categorization, and historical query filtering.")
            ]
            for r_idx, row in enumerate(modules_data):
                is_hdr = (r_idx == 0)
                for c_idx, val in enumerate(row):
                    format_cell(t.cell(r_idx, c_idx), val, is_header=is_hdr, font_size=11)
            print("Updated Slide 6 Module table")

    # ----------------------------------------------------
    # SLIDE 7 (06 Tech Stack): Add PDF and icon tools
    # ----------------------------------------------------
    slide7 = prs.slides[6]
    for s in slide7.shapes:
        if s.name == "Text 18" and s.has_text_frame:
            tools_items = [
                "jsPDF & jsPDF-AutoTable (client-side multi-page PDF generation with auto table formatting)",
                "Lucide React (modern dashboard & editor icons, visual metric indicators)",
                "AI Query Diagnostic Engine (heuristic rule-based parser & AST rewrite advisor)",
                "JWT & bcryptjs (token validation & secure credential hashing)",
                "Vite (lightning-fast frontend tooling, HMR & build optimization)",
                "MongoDB Node.js Driver (direct connection pool & compound index querying)"
            ]
            set_bullet_list(s.text_frame, tools_items, size=11.5)

    # ----------------------------------------------------
    # SLIDE 8 (07 UI Design / Wireframes): Insert UI Screen & Details
    # ----------------------------------------------------
    slide8 = prs.slides[7]
    for s in slide8.shapes:
        if s.name == "Text 8" and s.has_text_frame:
            s.text_frame.text = "UI Design & Core Screens"
            p = s.text_frame.paragraphs[0]
            p.font.name = "Calibri"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEAL_HEADER
        elif s.name == "Text 9" and s.has_text_frame:
            ui_items = [
                "Interactive Screen Architecture:",
                "• Authentication Portal: Tabbed sign-up, login, and JWT session handling",
                "• SQL Query Workbench: Code editor with sample query presets & validation",
                "• Real-Time Performance HUD: Score gauge (0-100), risk badge (Low/Med/High)",
                "• AI Advisor Panel: Actionable index recommendations & sargable rewrites",
                "• Analytics Dashboard: Historical query trends, average scores & report search",
                "• PDF Export Engine: Instant client-side download of DBA-ready performance audits",
                "Live Vercel Deployment: https://client-black-pi-23.vercel.app"
            ]
            set_bullet_list(s.text_frame, ui_items, size=11)

    slide8_img = r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\processed_screenshots\slide8_workbench.png"
    if os.path.exists(slide8_img):
        slide8.shapes.add_picture(slide8_img, Inches(1.25), Inches(1.95), Inches(6.50), Inches(4.40))
        print("Inserted slide8_workbench.png into Slide 8")

    # ----------------------------------------------------
    # SLIDE 9 (08 Frontend Development)
    # ----------------------------------------------------
    slide9 = prs.slides[8]
    for s in slide9.shapes:
        if s.name == "Text 5" and s.has_text_frame:
            fe_items = [
                "Built with React.js and Vite using modern functional components and custom React Hooks (useState, useEffect).",
                "Real-time query workbench featuring code editors, risk scoring gauges, and bottleneck telemetry.",
                "Modular components: Navigation Sidebar, Query Editor, Findings HUD, Index Advisor, and Analytics Dashboard.",
                "Client-side multi-page PDF generation via jspdf and jspdf-autotable with dynamic page-break calculations.",
                "Responsive dark-themed modern UI designed with Tailwind CSS principles and lucide-react iconography.",
                "Seamless backend communication using Fetch API with centralized JWT Bearer token authorization headers."
            ]
            set_bullet_list(s.text_frame, fe_items, size=12)

    # ----------------------------------------------------
    # SLIDE 10 (09 Backend Development)
    # ----------------------------------------------------
    slide10 = prs.slides[9]
    for s in slide10.shapes:
        if s.name == "Text 5" and s.has_text_frame:
            be_items = [
                "Node.js & Express.js RESTful API architecture with CORS, JSON body-parsing, and no-cache security middleware.",
                "Secure authentication pipeline utilizing bcryptjs for salted password hashing and jsonwebtoken for signed Bearer tokens.",
                "AI Performance Engine (services/analyzer.js): Static AST analysis detecting SELECT *, non-sargable functions (LOWER(), UPPER()), leading wildcards (LIKE '%term'), and unindexed joins.",
                "Automated query rewrite generator synthesizing covering index queries, anchored prefix searches, and bounded LIMIT clauses.",
                "Comprehensive endpoints: /api/auth/*, /api/analyze, /api/reports, /api/stats, and /api/health verifying MongoDB connectivity."
            ]
            set_bullet_list(s.text_frame, be_items, size=12)

    # ----------------------------------------------------
    # SLIDE 11 (10 Database Integration)
    # ----------------------------------------------------
    slide11 = prs.slides[10]
    for s in slide11.shapes:
        if s.name == "Text 5" and s.has_text_frame:
            db_items = [
                "Persistent MongoDB NoSQL database (query_performance_advisors) connected via official MongoDB Driver.",
                "users collection: Stores user credentials with an enforced unique lowercase index on email.",
                "reports collection: Stores diagnostic runs, raw SQL queries, cost scores, and AI recommendations.",
                "Compound index on { userId: 1, createdAt: -1 } for sub-second retrieval of user-scoped query history.",
                "Strict referential separation: Each user only accesses their own diagnostic reports via JWT identity verification.",
                "Health probe endpoint (/api/health) verifying real-time database connection state, latency, and active collection metadata."
            ]
            set_bullet_list(s.text_frame, db_items, size=12)

    # ----------------------------------------------------
    # SLIDE 12 (11 Real API Endpoints): Exact Replacement Table
    # ----------------------------------------------------
    slide12 = prs.slides[11]
    for s in slide12.shapes:
        if s.has_table:
            t = s.table
            api_data = [
                ("Method", "Endpoint", "Description"),
                ("POST", "/api/auth/register", "Register new user account with hashed password."),
                ("POST", "/api/auth/login", "Authenticate credentials and return JWT bearer token."),
                ("POST", "/api/analyze", "Submit query & execution plan to AI analyzer engine."),
                ("GET", "/api/reports", "Fetch all previously saved analysis reports for authenticated user."),
                ("GET", "/api/reports/:id", "Retrieve specific report details for inspection or PDF generation."),
                ("DELETE", "/api/reports/:id", "Remove a saved report from user query history.")
            ]
            for r_idx, row in enumerate(api_data):
                is_hdr = (r_idx == 0)
                for c_idx, val in enumerate(row):
                    format_cell(t.cell(r_idx, c_idx), val, is_header=is_hdr, font_size=11)
            print("Updated Slide 12 API table")

    # ----------------------------------------------------
    # SLIDE 13 (12 Implementation Status): QueryPilot Completion Breakdown
    # ----------------------------------------------------
    slide13 = prs.slides[12]
    # Update labels and progress bar widths
    status_specs = [
        # (label_shape_name, bar_shape_name, pct_shape_name, label_text, pct_val, pct_str)
        ("Text 5", "Shape 7", "Text 8", "User Authentication (JWT & bcrypt)", 1.00, "100%"),
        ("Text 9", "Shape 11", "Text 12", "Frontend UI (Query Editor, Metrics & Advisor View)", 0.90, "90%"),
        ("Text 13", "Shape 15", "Text 16", "Query Performance Analyzer Engine", 0.85, "85%"),
        ("Text 17", "Shape 19", "Text 20", "Database & Report Persistence (MongoDB)", 0.85, "85%"),
        ("Text 21", "Shape 23", "Text 24", "PDF Report Generation (jspdf-autotable)", 1.00, "100%"),
        ("Text 25", "Shape 27", "Text 28", "Deployment (Vercel Client + Cloud Backend)", 0.85, "85%"),
    ]
    max_bar_width = Inches(6.00)

    for l_name, b_name, p_name, l_text, pct, p_str in status_specs:
        for s in slide13.shapes:
            if s.name == l_name and s.has_text_frame:
                s.text_frame.text = l_text
                p = s.text_frame.paragraphs[0]
                p.font.name = "Calibri"
                p.font.size = Pt(12)
                p.font.color.rgb = COLOR_DARK_TEXT
            elif s.name == b_name:
                s.width = int(max_bar_width * pct)
                s.fill.solid()
                s.fill.fore_color.rgb = COLOR_TEAL_LIGHT if pct >= 0.9 else COLOR_TEAL_ACCENT
            elif s.name == p_name and s.has_text_frame:
                s.text_frame.text = p_str
                p = s.text_frame.paragraphs[0]
                p.font.name = "Calibri"
                p.font.size = Pt(12)
                p.font.bold = True
                p.font.color.rgb = COLOR_TEAL_HEADER
    print("Updated Slide 13 Implementation Status")

    # ----------------------------------------------------
    # SLIDE 14 (13 Screenshots): Insert 4 Live Screenshots
    # ----------------------------------------------------
    slide14 = prs.slides[13]
    # The 4 card coordinates:
    # 1: (1.20, 1.90, 5.30, 1.95)
    # 2: (6.80, 1.90, 5.30, 1.95)
    # 3: (1.20, 4.15, 5.30, 1.95)
    # 4: (6.80, 4.15, 5.30, 1.95)
    scr_files = [
        (r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\processed_screenshots\screen1_auth.png", 1.22, 1.92, 5.26, 1.91),
        (r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\processed_screenshots\screen2_workbench.png", 6.82, 1.92, 5.26, 1.91),
        (r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\processed_screenshots\screen3_telemetry.png", 1.22, 4.17, 5.26, 1.91),
        (r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\processed_screenshots\screen4_dashboard.png", 6.82, 4.17, 5.26, 1.91),
    ]
    for path, l, t, w, h in scr_files:
        if os.path.exists(path):
            slide14.shapes.add_picture(path, Inches(l), Inches(t), Inches(w), Inches(h))
    print("Inserted 4 screenshots into Slide 14")

    # ----------------------------------------------------
    # SLIDE 15 (14 Challenges Faced & Solutions): Exact Replacement
    # ----------------------------------------------------
    slide15 = prs.slides[14]
    for s in slide15.shapes:
        if s.has_table:
            t = s.table
            challenges_data = [
                ("Challenge", "Solution Adopted"),
                ("Complex SQL Parsing - Varied dialect syntaxes (PostgreSQL, MySQL)", 
                 "Implemented standardized heuristic tokenization and execution plan breakdown before passing to the advisor."),
                ("Multi-page PDF Formatting - Performance tables and code snippets were clipping", 
                 "Integrated jspdf-autotable with dynamic page-break calculations to ensure large queries and recommendation lists format cleanly."),
                ("Real-Time UI State Management Across Async Steps", 
                 "Centralized execution state using React hooks and modular async handler functions in App.jsx for smooth transitions across raw inputs, AST metrics, and AI recommendations."),
                ("Non-Sargable Query Detection Without DB Execution Overhead", 
                 "Developed static pattern matching for function wrappers (LOWER(), UPPER()) and wildcard scans (LIKE '%term') to flag index invalidation safely.")
            ]
            for r_idx, row in enumerate(challenges_data):
                is_hdr = (r_idx == 0)
                for c_idx, val in enumerate(row):
                    format_cell(t.cell(r_idx, c_idx), val, is_header=is_hdr, font_size=11)
            print("Updated Slide 15 Challenges table")

    # ----------------------------------------------------
    # SLIDE 16 (15 Next Steps / Remaining Work & Plan)
    # ----------------------------------------------------
    slide16 = prs.slides[15]
    for s in slide16.shapes:
        if s.name == "Text 5" and s.has_text_frame:
            next_steps = [
                "Native Execution Plan Ingestion: Direct JSON parsing for MySQL EXPLAIN FORMAT=JSON and PostgreSQL EXPLAIN (ANALYZE, BUFFERS).",
                "NoSQL / MongoDB Aggregation Optimization: Extend optimizer to detect unindexed $lookup stages, missing pipeline indexes, and memory-intensive $sort operations.",
                "Automated Migration Script Generator: Produce instant executable DDL scripts (CREATE INDEX ...) tailored to target database dialects.",
                "Historical Benchmark Comparison: Visual before-and-after query latency graphs tracking performance improvements post-indexing.",
                "Collaborative Team Workspaces: Shared team query repositories with role-based governance and audit history."
            ]
            set_bullet_list(s.text_frame, next_steps, size=12)
            print("Updated Slide 16 Next Steps")

    # ----------------------------------------------------
    # SLIDE 17 (16 References): Clean References
    # ----------------------------------------------------
    slide17 = prs.slides[16]
    for s in slide17.shapes:
        if s.has_table:
            t = s.table
            refs = [
                ("#", "Reference"),
                ("1", "MySQL 8.4 Reference Manual — EXPLAIN Statement & Query Execution Plan Optimization"),
                ("2", "PostgreSQL Documentation — Performance Tips, Indexes, and EXPLAIN Cost Modeling"),
                ("3", "MongoDB Official Documentation — Query Optimization, Aggregation Pipeline, and Index Strategies"),
                ("4", "Express.js & Node.js Documentation — REST API Design, Security Middleware, and JWT Authentication")
            ]
            for r_idx, row in enumerate(refs):
                is_hdr = (r_idx == 0)
                for c_idx, val in enumerate(row):
                    format_cell(t.cell(r_idx, c_idx), val, is_header=is_hdr, font_size=11)
            print("Updated Slide 17 References table")

    # Save to both target locations
    dst_downloads = r"C:\Users\HP\Downloads\QueryPilot_Project_Review_Presentation.pptx"
    dst_workspace = r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\QueryPilot_Project_Review_Presentation.pptx"
    
    prs.save(dst_downloads)
    prs.save(dst_workspace)
    print("SUCCESS: Presentation successfully saved to:")
    print(" -", dst_downloads)
    print(" -", dst_workspace)

if __name__ == "__main__":
    update_presentation()
