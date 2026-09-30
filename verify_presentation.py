import pptx

prs = pptx.Presentation(r"D:\Database\AI-Powered-Database-Query-Performance-Advisor-MERN\QueryPilot_Project_Review_Presentation.pptx")

print(f"Total slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides):
    shapes_info = []
    for s in slide.shapes:
        info = f"{s.name}({s.shape_type})"
        if s.has_text_frame:
            t = " ".join([p.text.strip() for p in s.text_frame.paragraphs if p.text.strip()])
            if t:
                info += f": '{t[:40]}...'"
        elif s.has_table:
            info += f": TABLE {len(s.table.rows)}x{len(s.table.columns)}"
        shapes_info.append(info)
    print(f"\n--- SLIDE {idx+1} ({len(slide.shapes)} shapes) ---")
    for si in shapes_info[:6]:
        print(" ", si)
